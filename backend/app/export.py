"""Write computed crosstabs to a plain-grid Excel workbook (one sheet per table).

Every table is driven by a saved spec, recomputed here so filters, weighting,
nested banners, groups, and display-only hides/renames all apply. Each cell
statistic gets its own row (one value per cell); significance marks are appended
to the column-percent (or first shown) statistic.
"""

from __future__ import annotations

import io

import pandas as pd

from . import crosstab as crosstabbing
from .models import (
    CrosstabRequest,
    CrosstabResponse,
    DatasetMeta,
    ExportTable,
    SavedCrosstabSpec,
)

CATSEP = "\u0001"

CELL_STATS = [
    ("count", "Count"),
    ("col_pct", "Column %"),
    ("row_pct", "Row %"),
    ("total_pct", "Total %"),
]
SUMMARY_ROWS = [
    ("base_n", "Base n"),
    ("eff_n", "Effective n"),
    ("total_count", "Column n"),
    ("total_sum", "Total Sum"),
    ("mean", "Mean"),
]
NUMERIC_SUMMARY = {"total_sum", "mean"}

_INVALID_SHEET = set(r"[]:*?/\\")


def _safe_sheet_name(name: str, used: set[str]) -> str:
    clean = "".join(" " if ch in _INVALID_SHEET else ch for ch in name).strip()
    clean = (clean or "Table")[:31]
    base, n = clean, 2
    while clean.lower() in used:
        suffix = f" ({n})"
        clean = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(clean.lower())
    return clean


def _request_from_spec(spec: SavedCrosstabSpec) -> CrosstabRequest:
    return CrosstabRequest(
        row=spec.row,
        column=spec.column,
        banner=spec.banner,
        banner_groups=spec.banner_groups,
        filter_id=spec.filter_id,
        weight=spec.weight,
        row_groups=spec.row_groups,
        column_groups=spec.column_groups,
    )


def build_workbook(
    meta: DatasetMeta, df: pd.DataFrame, tables: list[ExportTable]
) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for i, table in enumerate(tables):
        ws = wb.create_sheet(title=_safe_sheet_name(table.name or f"Table {i + 1}", used))
        try:
            result = crosstabbing.compute_crosstab(df, meta, _request_from_spec(table.spec))
            _render_sheet(ws, meta, table.spec, result)
        except crosstabbing.CrosstabError as err:
            ws["A1"] = f"Could not build this table: {err}"
    if not wb.sheetnames:
        wb.create_sheet(title="Empty")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _render_sheet(
    ws, meta: DatasetMeta, spec: SavedCrosstabSpec, result: CrosstabResponse
) -> None:
    advanced = bool(spec.banner)
    display = spec.display
    cell_stats = [(k, lbl) for k, lbl in CELL_STATS if k in display.cell_stats]
    numeric_ok = any(v is not None for v in result.row_values)
    summary_rows = [
        (k, lbl)
        for k, lbl in SUMMARY_ROWS
        if k in display.summary_rows and (numeric_ok or k not in NUMERIC_SUMMARY)
    ]
    letters_on = "letters" in display.significance
    arrows_on = "arrows" in display.significance

    def cat_key(c) -> str:
        return f"{c.seg}{CATSEP}{c.label}"

    vis_cols: list[int] = []
    for ci, c in enumerate(result.columns):
        if advanced:
            if cat_key(c) in spec.banner_cat_hidden:
                continue
            if (c.group or "") in spec.banner_parent_hidden:
                continue
        elif c.label in spec.column_hidden:
            continue
        vis_cols.append(ci)
    vis_rows = [
        ri for ri, lbl in enumerate(result.row_labels) if lbl not in spec.row_hidden
    ]
    # Display-only reorder: labels in spec order first, others keep natural order.
    # Columns reorder leaves only within their banner parent/segment (same group).
    if spec.row_order:
        rank = {lbl: i for i, lbl in enumerate(spec.row_order)}
        vis_rows.sort(key=lambda ri: (rank.get(result.row_labels[ri], len(rank)), ri))
    if spec.column_order:
        rank = {lbl: i for i, lbl in enumerate(spec.column_order)}
        runs: list[list[int]] = []
        for ci in vis_cols:
            g = result.columns[ci].group or ""
            if runs and (result.columns[runs[-1][0]].group or "") == g:
                runs[-1].append(ci)
            else:
                runs.append([ci])
        ordered: list[int] = []
        for run in runs:
            run.sort(key=lambda ci: (rank.get(result.columns[ci].label, len(rank)), ci))
            ordered.extend(run)
        vis_cols = ordered
    sig_on = (letters_on or arrows_on) and len(vis_cols) >= 2

    def disp_col(c) -> str:
        if advanced:
            return spec.banner_cat_renames.get(cat_key(c), c.label)
        return spec.column_renames.get(c.label, c.label)

    def disp_top(c) -> str:
        return spec.banner_parent_renames.get(c.group or "", c.top_label or "")

    def disp_row(lbl: str) -> str:
        return spec.row_renames.get(lbl, lbl)

    two_level = advanced and any(result.columns[ci].top_label for ci in vis_cols)

    # Margins (mirror the frontend: totals span every column / row).
    n_rows = len(result.cells)
    row_totals = [sum(cell.count for cell in row) for row in result.cells]
    col_sums: list[float] = []
    col_valid: list[float] = []
    for ci in range(len(result.columns)):
        s = 0.0
        v = 0.0
        for ri in range(n_rows):
            value = result.row_values[ri]
            if value is None:
                continue
            s += value * result.cells[ri][ci].count
            v += result.cells[ri][ci].count
        col_sums.append(s)
        col_valid.append(v)
    col_totals = [
        sum(result.cells[ri][ci].count for ri in range(n_rows))
        for ci in range(len(result.columns))
    ]

    def stat_value(ri: int, ci: int, key: str) -> float | None:
        count = result.cells[ri][ci].count
        base = result.columns[ci].base
        if key == "count":
            return count
        if key == "col_pct":
            return count / base * 100.0 if base else None
        if key == "row_pct":
            return count / row_totals[ri] * 100.0 if row_totals[ri] else None
        if key == "total_pct":
            return (
                count / result.total_base * 100.0 if result.total_base else None
            )
        return None

    def summary_value(key: str, ci: int) -> float | None:
        col = result.columns[ci]
        if key == "base_n":
            return col.base
        if key == "eff_n":
            return col.eff_base if col.eff_base is not None else col.base
        if key == "total_count":
            return col_totals[ci]
        if key == "total_sum":
            return col_sums[ci]
        if key == "mean":
            return col_sums[ci] / col_valid[ci] if col_valid[ci] else None
        return None

    def marks(ri: int, ci: int) -> str:
        cell = result.cells[ri][ci]
        parts = ""
        if arrows_on and cell.sig_arrow == "up":
            parts += "\u25b2"
        elif arrows_on and cell.sig_arrow == "down":
            parts += "\u25bc"
        if letters_on and cell.sig_higher:
            parts += (" " if parts else "") + " ".join(cell.sig_higher)
        return parts

    attach_stat = None
    if cell_stats:
        keys = [k for k, _ in cell_stats]
        attach_stat = "col_pct" if "col_pct" in keys else keys[0]

    first_col = 3  # A=row label, B=statistic, data from column C
    r = 1

    if two_level:
        prev_group = object()
        pos = first_col
        for ci in vis_cols:
            c = result.columns[ci]
            if c.group != prev_group:
                ws.cell(r, pos, disp_top(c))
                prev_group = c.group
            pos += 1
        r += 1

    pos = first_col
    for ci in vis_cols:
        c = result.columns[ci]
        label = disp_col(c)
        if letters_on and sig_on and c.letter:
            label = f"{label} ({c.letter})"
        ws.cell(r, pos, label)
        pos += 1
    r += 1

    for ri in vis_rows:
        for key, label in cell_stats:
            ws.cell(r, 1, disp_row(result.row_labels[ri]))
            ws.cell(r, 2, label)
            pos = first_col
            for ci in vis_cols:
                value = stat_value(ri, ci, key)
                cell = ws.cell(r, pos)
                mark = marks(ri, ci) if sig_on and key == attach_stat else ""
                if value is None:
                    cell.value = mark or ""
                elif key == "count":
                    if mark:
                        cell.value = f"{round(value)} {mark}"
                    else:
                        cell.value = round(value)
                else:
                    if mark:
                        cell.value = f"{value:.1f}% {mark}"
                    else:
                        cell.value = round(value, 1)
                        cell.number_format = '0.0"%"'
                pos += 1
            r += 1

    for key, label in summary_rows:
        ws.cell(r, 1, label)
        pos = first_col
        for ci in vis_cols:
            value = summary_value(key, ci)
            cell = ws.cell(r, pos)
            if value is None:
                cell.value = ""
            elif key in NUMERIC_SUMMARY:
                cell.value = round(value, 2)
            else:
                cell.value = round(value)
            pos += 1
        r += 1

    r += 1
    ws.cell(r, 1, _caption(meta, spec, result))
    r += 1
    if sig_on:
        ws.cell(r, 1, _legend(arrows_on, letters_on))

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 12
    for i in range(len(vis_cols)):
        ws.column_dimensions[_col_letter(first_col + i)].width = 14


def _caption(meta: DatasetMeta, spec: SavedCrosstabSpec, result: CrosstabResponse) -> str:
    total = round(result.total_base)
    if not result.weighted:
        return f"Unweighted, Total sample = {total}"
    label = spec.weight or ""
    if spec.weight:
        var = next((v for v in meta.variables if v.name == spec.weight), None)
        label = var.label if var else spec.weight
    text = f"Weighted \u2013 {label}, Total sample = {total}"
    if result.total_eff_base is not None:
        eff = round(result.total_eff_base)
        efficiency = (
            result.total_eff_base / result.total_base * 100.0
            if result.total_base
            else 0.0
        )
        text += f", Effective sample = {eff}, Weighting efficiency = {efficiency:.1f}%"
    return text


def _legend(arrows_on: bool, letters_on: bool) -> str:
    parts = []
    if arrows_on:
        parts.append("\u25b2/\u25bc = higher/lower than the rest of the sample")
    if letters_on:
        parts.append("letters = columns this cell is significantly higher than")
    return "Significance at 95%: " + "; ".join(parts) + "."


def _col_letter(idx: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(idx)
